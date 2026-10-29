"""Weekly report of 10-14, built on the lab's weekly report file (title slide + one content slide).
Story: the capstone stacking system rebuilt in simulation and run in various environments (picture: fig_weekly_envs.py)."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SRC, OUT = Path(sys.argv[1]), HERE / "주간보고_10_14_이지용.pptx"
DATE, NAME = "2026. 10. 14", "Jiyong Lee"

PREVIOUS = [
    (1, "Capstone project (smart factory stacking) in simulation"),
    (2, "I rebuilt our capstone system in MuJoCo because the robot arm was broken"),
    (3, "It used the same SO-101 arm, scale, blue bins and cameras as the capstone"),
    (2, "I trained ACT, and the robot stacked three bins with 99.0% success"),
    (2, "I ran it in various environments, including lighting changes (88–100%)"),
]
NEXT = [
    (1, "I will submit the paper for the undergraduate paper competition by Oct. 19"),
    (1, "I will submit the RISE support documents"),
]
SIZE = {0: 20, 1: 15, 2: 13, 3: 11.5}


def set_text(shape, text):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def fill(shape, head, items):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = head; r.font.size = Pt(SIZE[0]); r.font.bold = True
    for lvl, text in items:
        p = tf.add_paragraph(); p.level = lvl
        r = p.add_run(); r.text = text; r.font.size = Pt(SIZE[lvl])


prs = Presentation(SRC)
s1, s2 = prs.slides[0], prs.slides[1]
set_text(next(sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("2026")), DATE)
# The text boxes of the reference sit in mc:AlternateContent (math for "π0.5"); keep the plain p:sp of the Choice part
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
for ac in list(s2.shapes._spTree.iter(MC + "AlternateContent")):
    sp = ac.find(MC + "Choice").find("{http://schemas.openxmlformats.org/presentationml/2006/main}sp")
    ac.addprevious(sp)
    ac.getparent().remove(ac)
shapes = list(s2.shapes)
set_text(next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Seokhyeon Shin"), NAME)
prev = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("Previous"))
nxt = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("Next"))
for sh in shapes:  # the reference picture
    if sh.shape_type == 13:
        sh._element.getparent().remove(sh._element)
fill(prev, "Previous", PREVIOUS)
fill(nxt, "Next", NEXT)
W = prs.slide_width
prev.left, prev.top, prev.width, prev.height = Inches(0.25), Inches(0.85), W - Inches(0.45), Inches(2.3)
pic_w = W - Inches(0.5)
pic = s2.shapes.add_picture(str(HERE / "weekly_1014_envs.png"), Inches(2.1), Inches(2.75), width=Inches(5.8))
nxt.left, nxt.top, nxt.width, nxt.height = Inches(0.25), pic.top + pic.height + Inches(0.02), W - Inches(0.45), Inches(1.0)
# The removed reference picture and the fallback pictures of its text boxes stay in the file through the
# slide's relationships; drop every picture the slide no longer shows
used = set(s2._element.xpath(".//@r:embed | .//@r:id | .//@r:link"))
for rId in [k for k, rel in s2.part.rels.items() if rel.reltype.endswith("/image") and k not in used]:
    s2.part.drop_rel(rId)
prs.core_properties.title = "Weekly Report " + DATE
prs.core_properties.author = ""
prs.core_properties.last_modified_by = ""
prs.save(OUT)
# docProps/app.xml keeps the slide titles of the reference file ("TitlesOfParts"); put this report's name there
import shutil, tempfile, zipfile
tmp = Path(tempfile.mkdtemp()) / "x.pptx"
with zipfile.ZipFile(OUT) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        data = zi.read(it.filename)
        if it.filename == "docProps/app.xml":
            data = data.decode("utf8").replace("Seokhyeon Shin", NAME).encode("utf8")
        zo.writestr(it, data)
shutil.move(tmp, OUT)
print("wrote", OUT, "picture height %.2f in" % (pic.height / 914400))
