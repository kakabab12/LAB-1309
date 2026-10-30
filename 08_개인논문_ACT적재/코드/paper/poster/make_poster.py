"""Conference poster (IEIE 2026 Fall, poster session): one 900 x 1200 mm slide, editable in PowerPoint.

IEIE poster rules (conf.theieie.org/2026f, 논문발표 안내): panel 100 x 180 cm; title / authors / affiliation 90 x 20 cm;
letters: title and authors 25-30 mm, chapter titles 15-20 mm, text 7-10 mm; A4 x 12, B4 x 10 or one full sheet.
-> one sheet 90 x 120 cm: title band 900 x 200 mm, title 84 pt (29.6 mm), authors 72 pt (25.4 mm),
   section titles 50 pt (17.6 mm), text 24 pt (8.5 mm), notes 20 pt (7.1 mm).
Numbers are the paper's (simulation unless noted; weight recognition measured on the real device).
usage: PYTHONPATH=<python-pptx> python3 paper/poster/make_poster.py [out.pptx]
"""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "poster.pptx"

FONT = "맑은 고딕"
NAVY, INK, MUTED = "14213D", "1F2933", "52606D"
BLUE, ORANGE, AMBER = "1C5FC8", "D9480F", "FFC23D"
TINT, LINE, GREEN = "EDF2FA", "C9D3E3", "2B8A3E"
W, H = 900, 1200
COL_L, COL_R, COL_W = 22, 460, 418


def mm(v: float) -> Emu:
    return Emu(int(round(v * 36000)))


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = mm(W), mm(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SH = slide.shapes


# ------------------------------------------------------------------ text helpers
def style_run(r, size, bold=False, color=INK, italic=False, baseline=None):
    f = r.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)
    if baseline is not None:
        rpr.set("baseline", str(baseline))


def para_props(p, line=1.12, after=0.0, before=0.0, bullet=None, indent=0.0, align=None):
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag in (qn("a:lnSpc"), qn("a:spcBef"), qn("a:spcAft"), qn("a:buFont"), qn("a:buChar"), qn("a:buNone")):
            pPr.remove(child)
    els = []
    ln = etree.Element(qn("a:lnSpc")); etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(line * 100000))); els.append(ln)
    sb = etree.Element(qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts")).set("val", str(int(before * 100))); els.append(sb)
    sa = etree.Element(qn("a:spcAft")); etree.SubElement(sa, qn("a:spcPts")).set("val", str(int(after * 100))); els.append(sa)
    if bullet:
        bf = etree.Element(qn("a:buFont")); bf.set("typeface", "Arial"); els.append(bf)
        bc = etree.Element(qn("a:buChar")); bc.set("char", bullet); els.append(bc)
        pPr.set("marL", str(int(mm(indent)))); pPr.set("indent", str(-int(mm(indent))))
    else:
        els.append(etree.Element(qn("a:buNone")))
    for i, el in enumerate(els):
        pPr.insert(i, el)
    if align is not None:
        p.alignment = align


def textbox(x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, margin=0.0):
    """paras: list of dicts {runs: [(text, size, bold, color[, baseline])], bullet, line, after, align, indent}."""
    tb = SH.add_textbox(mm(x), mm(y), mm(w), mm(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", mm(margin))
    for k, pd in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        for spec in pd["runs"]:
            text, size, bold, color = spec[:4]
            r = p.add_run(); r.text = text
            style_run(r, size, bold, color, baseline=spec[4] if len(spec) > 4 else None)
        para_props(p, line=pd.get("line", 1.12), after=pd.get("after", 0), before=pd.get("before", 0),
                   bullet=pd.get("bullet"), indent=pd.get("indent", 6.5), align=pd.get("align"))
    return tb


def P(text, size=24, bold=False, color=INK, **kw):
    return {"runs": [(text, size, bold, color)], **kw}


def B(*runs, size=24, after=3.0):
    """Bullet paragraph from (text, bold?, color?) parts."""
    rr = []
    for part in runs:
        if isinstance(part, str):
            rr.append((part, size, False, INK))
        else:
            t, bold, col = (part + (INK,))[:3] if len(part) == 2 else part
            rr.append((t, size, bold, col))
    return {"runs": rr, "bullet": "•", "indent": 6.5, "after": after, "line": 1.1}


def box(x, y, w, h, fill, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    s = SH.add_shape(shape, mm(x), mm(y), mm(w), mm(h))
    if name:
        s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(2)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def shape_text(s, paras, anchor=MSO_ANCHOR.MIDDLE, margin=2.0):
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", mm(margin))
    for k, pd in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        for spec in pd["runs"]:
            r = p.add_run(); r.text = spec[0]; style_run(r, *spec[1:4])
        para_props(p, line=pd.get("line", 1.05), after=pd.get("after", 0), align=pd.get("align", PP_ALIGN.CENTER))


def picture(path, x, y, w=None, h=None, name=None):
    pic = SH.add_picture(str(path), mm(x), mm(y), mm(w) if w else None, mm(h) if h else None)
    if name:
        pic.name = name
    return pic


def label(x, y, text, size=22, w=None, color=INK, fill="FFFFFF"):
    """Small white tag over a picture."""
    w = w or (len(text) * size * 0.36 + 8)
    s = box(x, y, w, size * 0.36 + 6, fill, line="9AA5B1", radius=0.25)
    shape_text(s, [P(text, size, True, color, align=PP_ALIGN.CENTER)], margin=1.0)
    return s


def section(n, title, x, y, w=COL_W):
    c = box(x, y + 1, 22, 22, AMBER, shape=MSO_SHAPE.OVAL, name=f"badge {n}")
    shape_text(c, [P(str(n), 34, True, NAVY, align=PP_ALIGN.CENTER)], margin=0)
    textbox(x + 28, y - 1, w - 28, 25, [P(title, 50, True, NAVY, line=1.0)], anchor=MSO_ANCHOR.MIDDLE,
            name=f"section {n}")


def sub(text):
    return (text, 24, False, INK, -25000)


# ------------------------------------------------------------------ title band (900 x 200 mm)
band = box(0, 0, W, 200, NAVY, shape=MSO_SHAPE.RECTANGLE, name="title band")
textbox(30, 14, 840, 14, [P("2026 대한전자공학회 추계학술대회 · 학부생 논문경진대회", 26, True, AMBER)], name="conference")
textbox(30, 32, 840, 70, [P("ACT 모방학습 기반 스마트 팩토리 순차 적재에서", 84, True, "FFFFFF", line=1.0),
                          P("시연 일관성의 영향 분석", 84, True, "FFFFFF", line=1.0)], name="title")
textbox(30, 106, 840, 16, [P("Analyzing the Effects of Demonstration Consistency on ACT-Based Sequential Stacking "
                             "for Smart Factories", 30, False, "C9D6EA")], name="title en")
textbox(30, 130, 840, 40, [{"runs": [("이지용, 고병철", 72, True, "FFFFFF"),
                                     ("     계명대학교 컴퓨터공학전공", 40, False, "D6E0F0")], "line": 1.0}],
        anchor=MSO_ANCHOR.MIDDLE, name="authors")

# ------------------------------------------------------------------ left column
y = 220
section(1, "연구 배경 및 목표", COL_L, y)
textbox(COL_L, y + 30, COL_W, 52, [
    B("중소 공정은 고가 산업용 로봇 도입이 어려움 → 저가 로봇팔과 모방학습(ACT", ("[1]", False, MUTED), ")으로 적은 시연에서 조작을 학습"),
    B("순차 적재는 앞 단계가 놓은 적재함이 다음 단계의 출발점 → ", ("작은 오차와 서로 다른 시연 동작이 단계마다 누적", True, ORANGE)),
    B(("목표", True, NAVY), ": 모델 구조·학습 설정은 그대로 두고 ", ("시연의 일관성", True, BLUE), "만으로 세 단계 연속 성공률을 높인다", after=0),
], name="background")

y = 312
section(2, "캡스톤 시스템 구성", COL_L, y)
steps = [(["색 분류", "컨베이어"], ["파랑 큐브만"]), (["저울 숫자 판독"], ["YOLOv8n ·", "Jetson Orin Nano"]),
         (["118\u00a0g 이상"], ["TCP 트리거"]), (["SO-101 로봇팔"], ["단계별 ACT"]), (["1층 → 옆 → 2층"], ["순차 적재"])]
bw, gap = 70.0, 17.0
for i, (head, tail) in enumerate(steps):
    bx = COL_L + i * (bw + gap)
    s = box(bx, y + 32, bw, 52, TINT if i != 3 else "DCE7F8", line=LINE, radius=0.12, name=f"flow {i + 1}")
    shape_text(s, [P(h_, 22, True, NAVY, line=1.0) for h_ in head] + [P(t_, 18, False, MUTED, line=1.0) for t_ in tail],
               margin=1.2)
    if i < len(steps) - 1:
        a = SH.add_shape(MSO_SHAPE.RIGHT_ARROW, mm(bx + bw + 2.5), mm(y + 52), mm(gap - 5), mm(12))
        a.fill.solid(); a.fill.fore_color.rgb = rgb(NAVY); a.line.fill.background(); a.shadow.inherit = False
loop = SH.add_shape(MSO_SHAPE.U_TURN_ARROW, mm(COL_L + 3 * (bw + gap) + 8), mm(y + 86), mm(bw + gap + bw - 16), mm(14))
loop.rotation = 180
loop.fill.solid(); loop.fill.fore_color.rgb = rgb(AMBER); loop.line.fill.background(); loop.shadow.inherit = False
textbox(COL_L, y + 88, 3 * (bw + gap), 12, [P("무게 확인 재시도: 저울에 60 g 이상 남으면 그 단계를 다시 실행 →", 20, True, "8A5A00",
                                            align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE, name="retry note")
stats = [("86.3%", "숫자 탐지 mAP50"), ("91.7%", "118 g 분류 정확도 (55/60)"), ("약 810 ms", "인식 → 로봇 명령")]
sw = (COL_W - 2 * 8) / 3
for i, (num, cap) in enumerate(stats):
    sx = COL_L + i * (sw + 8)
    s = box(sx, y + 108, sw, 46, "FFFFFF", line=LINE, radius=0.1, name=f"stat {i + 1}")
    shape_text(s, [P(num, 44, True, BLUE, line=1.0), P(cap, 19, False, MUTED, line=1.0)], margin=1.0)
textbox(COL_L, y + 156, COL_W, 10, [P("무게 인식은 실제 장치에서 측정 (조건별 20회, 총 60회)", 20, False, MUTED)], name="real note")

y = 492
section(3, "시뮬레이션으로 재현 (MuJoCo)", COL_L, y)
hero_w, cam_w = 302.7, 109.0
hero_h = hero_w * 900 / 1600
picture(IMG / "p_hero.png", COL_L, y + 31, w=hero_w, name="sim hero")
cam_h = cam_w * 600 / 800
picture(IMG / "basic_cam_front.png", COL_L + hero_w + 6.3, y + 31, w=cam_w, name="wrist cam")
picture(IMG / "basic_cam_top.png", COL_L + hero_w + 6.3, y + 31 + cam_h + (hero_h - 2 * cam_h), w=cam_w, name="top cam")
label(COL_L + 8, y + 31 + hero_h - 22, "저울 (집는 곳)")
label(COL_L + 226, y + 31 + 112, "팔레트 (쌓는 곳)")
label(COL_L + hero_w + 9, y + 34, "손목 카메라", size=20)
label(COL_L + hero_w + 9, y + 31 + hero_h - cam_h + 3, "상단 카메라", size=20)
textbox(COL_L, y + 31 + hero_h + 5, COL_W, 50, [
    B("로봇팔이 고장 나 ", ("캡스톤 시스템을 그대로 재현", True, NAVY), ": 공식 SO-101 모델, 같은 카메라 위치·30 Hz 제어, 적재함 64×64×52 mm·123 g", size=22, after=2),
    B("적재함 위치 ±20 mm·각도 ±20° 무작위, 단계마다 스크립트 시연 100회로 ACT 학습 (LeRobot", ("[2]", False, MUTED), ")", size=22, after=0),
], name="sim notes")

y = 756
section(4, "제안 1. 시연 설계 규칙 3가지", COL_L, y)
rules = [("① 여유 두고 잡기", "벽에 바짝 붙여 내려가면 3\u00a0mm만 어긋나도 벽 위에 얹힘 → 벽 바깥 9\u00a0mm에서 내려감", "28% → 62%"),
         ("② 늘 같은 곳 잡기", "가장 가까운 벽을 잡으면 +7°·+9°처럼 비슷한 장면에서 잡는 벽이 바뀜 → 항상 같은 벽", "62% → 82%"),
         ("③ 놓고 물러나기", "놓자마자 올라가면 집게가 적재함을 끌어 올려 기울임 → 5\u00a0mm 물러난 뒤 올라감", "82% → 96%")]
from PIL import Image as _Image
rw, rh = 118.0, 80.0  # each schematic fitted inside this box (before = left half, after = right half)
for i, (head, desc, eff) in enumerate(rules):
    ry = y + 31 + i * (rh + 4.5)
    iw, ih = _Image.open(IMG / f"p_rule{i + 1}.png").size
    k = min(rw / iw, (rh - 4) / ih)
    picture(IMG / f"p_rule{i + 1}.png", COL_L + (rw - iw * k) / 2, ry + (rh - ih * k) / 2, w=iw * k, name=f"rule {i + 1}")
    textbox(COL_L + rw + 8, ry, COL_W - rw - 8, rh, [
        P(head, 30, True, NAVY, line=1.0, after=3),
        P(desc, 22, False, INK, line=1.1, after=3),
        {"runs": [("연속 성공률 ", 22, False, MUTED), (eff, 26, True, BLUE)], "line": 1.0},
    ], anchor=MSO_ANCHOR.MIDDLE, name=f"rule text {i + 1}")
    if i < len(rules) - 1:
        ln = SH.add_connector(1, mm(COL_L), mm(ry + rh + 2.25), mm(COL_L + COL_W), mm(ry + rh + 2.25))
        ln.line.color.rgb = rgb(LINE); ln.line.width = Pt(1.5)

# ------------------------------------------------------------------ right column
y = 220
section(5, "제안 2. 학습 전 시연 일관성 점검", COL_R, y)
textbox(COL_R, y + 30, COL_W, 44, [
    P("그리퍼가 닫히는 순간의 손목 회전각을 시연마다 모아 작은 값부터 정렬하고, 이웃한 값 사이의 최대 간격 G를 잰다", 22, line=1.08, after=3),
    # subscripts: PowerPoint draws baseline-shifted text at about 2/3 of its size, so they are set at 28 pt
    {"runs": [("G = max", 28, True, NAVY), ("1≤i<N", 28, False, NAVY, -25000), (" (θ", 28, True, NAVY), ("i+1", 28, False, NAVY, -25000),
              (" − θ", 28, True, NAVY), ("i", 28, False, NAVY, -25000), (")", 28, True, NAVY),
              ("     잡는 방식이 둘로 갈리면 큰 빈 구간 → 학습 없이 시연만으로 점검", 21, False, ORANGE)], "line": 1.0},
], name="audit text")
aud_w = COL_W
aud_h = aud_w * 215 / 400
picture(IMG / "audit.png", COL_R, y + 68, w=aud_w, name="audit figure")

y = 220 + 68 + aud_h + 14
section(6, "실험 결과 (시뮬레이션)", COL_R, y)
callouts = [("28% → 99.0%", "세 단계 연속 성공률 (시연 100회)", BLUE), ("200/200", "무게 확인 재시도까지 적용", NAVY)]
cw = (COL_W - 8) / 2
for i, (num, cap, col) in enumerate(callouts):
    s = box(COL_R + i * (cw + 8), y + 30, cw, 46, TINT, radius=0.12, name=f"callout {i + 1}")
    shape_text(s, [P(num, 50, True, col, line=1.0), P(cap, 20, False, MUTED, line=1.0)], margin=1.0)


def column_chart(x, y0, w, h, cats, series, colors, labels, legend=False, name=None):
    cd = CategoryChartData()
    cd.categories = cats
    for sname, vals in series:
        cd.add_series(sname, vals)
    gf = SH.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, mm(x), mm(y0), mm(w), mm(h), cd)
    if name:
        gf.name = name
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(18), FONT
    ch.font.color.rgb = rgb(INK)
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(18)
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = 112, 0
    va.has_major_gridlines = False
    va.visible = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(18)
    ca.format.line.color.rgb = rgb("9AA5B1")
    ca.has_major_gridlines = False
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.overlap = -10 if len(series) > 1 else 0
    for si, s in enumerate(plot.series):
        if not isinstance(colors[si], list):  # series fill too, so the legend shows the bar colour
            s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(colors[si])
        for pi, pt in enumerate(s.points):
            col = colors[si][pi] if isinstance(colors[si], list) else colors[si]
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
            dl = pt.data_label
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
            tf = dl.text_frame
            tf.text = labels[si][pi]
            r = tf.paragraphs[0].runs[0]
            r.font.size, r.font.bold = Pt(22), True
            r.font.color.rgb = rgb(col if col in (BLUE, NAVY) else INK)
            r.font.name = FONT
    cs = ch._chartSpace  # no frame around the chart area
    sp = cs.find(qn("c:spPr"))
    if sp is None:
        sp = etree.Element(qn("c:spPr"))
        cs.find(qn("c:chart")).addnext(sp)
    for el in list(sp):
        sp.remove(el)
    etree.SubElement(sp, qn("a:noFill"))
    etree.SubElement(etree.SubElement(sp, qn("a:ln")), qn("a:noFill"))
    return gf


cy = y + 82
textbox(COL_R, cy, 262, 10, [P("규칙을 하나씩 추가 (단계당 시연 100회)", 21, True, NAVY)], name="chart a title")
column_chart(COL_R - 2, cy + 9, 266, 128,
             ["초기", "+ 여유\n두고 잡기", "+ 늘 같은\n곳 잡기", "+ 놓고\n물러나기", "규칙 3개\n(200개 장면)"],
             [("연속 성공률", (28, 62, 82, 96, 99.0))], [["ADB5BD", "8DB0E8", "4F86D9", BLUE, NAVY]],
             [["28%", "62%", "82%", "96%", "99.0%"]], name="chart rules")
textbox(COL_R + 270, cy, COL_W - 270, 10, [P("시연 수를 늘리면?", 21, True, NAVY)], name="chart b title")
column_chart(COL_R + 266, cy + 9, COL_W - 264, 128, ["여유만\n(일관성 X)", "규칙\n시연"],
             [("시연 100회", (62, 96)), ("200회·1,000회", (36, 94))], [BLUE, "9AA5B1"], [["62%", "96%"], ["36%", "94%"]],
             legend=True, name="chart data")
textbox(COL_R, cy + 140, COL_W, 20, [
    {"runs": [("시연을 늘려도 일관성이 없으면 오히려 하락 (200회 36%) → ", 21, False, INK), ("양보다 일관성", 22, True, ORANGE)],
     "line": 1.0},
], name="result note")

y = cy + 140 + 22
section(7, "다양한 환경으로 확장", COL_R, y)
tiles = [("p_clutter.png", "잡동사니 책상", "92%"), ("p_factory.png", "공장형 장면", "90%"), ("p_looks.png", "조명·책상 변화", "98%*"),
         ("p_cup.png", "손잡이 컵", "88%"), ("p_box.png", "직사각형 상자", "100%"), ("p_pallet.png", "8칸 팔레트", "95%")]
tw = (COL_W - 2 * 6) / 3
th = tw * 836 / 1400
for i, (f, name, val) in enumerate(tiles):
    tx = COL_R + (i % 3) * (tw + 6)
    ty = y + 31 + (i // 3) * (th + 21)
    if (IMG / f).exists():
        picture(IMG / f, tx, ty, w=tw, name=f"tile {name}")
    else:
        s = box(tx, ty, tw, th, "E5E9F0", shape=MSO_SHAPE.RECTANGLE)
        shape_text(s, [P("[그림 준비 중]", 20, False, MUTED)])
    textbox(tx, ty + th + 1, tw, 16, [{"runs": [(name, 21, True, NAVY), ("  " + val, 24, True, BLUE)], "line": 1.0}],
            anchor=MSO_ANCHOR.MIDDLE, name=f"tile label {i + 1}")
textbox(COL_R, y + 31 + 2 * (th + 21) - 3, COL_W, 16, [
    P("세 단계 연속 성공률 (팔레트: 8칸을 모두 쌓은 비율 38/40). *15개 조건의 단계별 평균. 각 환경의 시연을 더해 학습", 20, False, MUTED,
      line=1.0)], name="tile note")

# ------------------------------------------------------------------ conclusion + footer (full width)
cyb = 1046
c = box(COL_L, cyb, W - 2 * COL_L, 76, TINT, radius=0.08, name="conclusion box")
section(8, "결론 및 향후 연구", COL_L + 6, cyb + 26, w=300)
textbox(COL_L + 300, cyb + 6, W - 2 * COL_L - 306, 66, [
    B("모델은 그대로 두고 시연 규칙만 바꿔 세 단계 연속 성공률 ", ("28% → 99.0%", True, BLUE), ", 무게 확인 재시도로 ", ("200/200", True, NAVY), size=22, after=2),
    B("학습 전 손목 각도 간격 G로 상충하는 시연을 찾음 → ", ("시연은 양보다 일관성", True, ORANGE), size=22, after=2),
    B("한계: 로봇 조작은 시뮬레이션·스크립트 시연 결과 → 실제 로봇과 사람 원격조작 시연으로 검증 예정", size=22, after=0),
], anchor=MSO_ANCHOR.MIDDLE, name="conclusion")
textbox(COL_L, 1126, 540, 64, [
    P("참고문헌", 20, True, NAVY, line=1.0, after=1),
    P("[1] T. Z. Zhao et al., “Learning Fine-Grained Bimanual Manipulation with Low-Cost Hardware,” RSS, 2023.", 20, False, MUTED, line=1.0),
    P("[2] R. Cadene et al., “LeRobot,” 2024, github.com/huggingface/lerobot", 20, False, MUTED, line=1.0),
    P("[3] S. Belkhale, Y. Cui, D. Sadigh, “Data Quality in Imitation Learning,” NeurIPS, 2023.", 20, False, MUTED, line=1.0),
    P("[4] E. Todorov, T. Erez, Y. Tassa, “MuJoCo: A Physics Engine for Model-Based Control,” IROS, 2012.", 20, False, MUTED, line=1.0),
], name="references")
textbox(574, 1126, W - COL_L - 574, 64, [
    P("본 과제(결과물)는 2026년도 교육부 및 대구광역시의 재원으로 RISE센터의 지원을 받아 수행된 지역혁신중심 대학지원체계(RISE)의 "
      "결과입니다.(과제관리번호: 2026-RISE-03-002)", 20, False, MUTED, line=1.05),
], anchor=MSO_ANCHOR.BOTTOM, name="acknowledgement")

prs.core_properties.author = ""
prs.core_properties.title = "ACT 모방학습 기반 스마트 팩토리 순차 적재에서 시연 일관성의 영향 분석 (포스터)"
prs.save(OUT)
print("wrote", OUT)
